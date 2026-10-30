"""Build Mercury-style part numbers from a dev export and check which exist in a Mercury export."""
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_SCAN_ROWS = 20
THREE_DIGIT_TAIL = re.compile(r'^(.+-)\d(\d{2})$')
CHECKED_PREFIXES = ('ASY-', 'PCA-')  # IC and other part types are skipped
NO_NEED_PROJECT_WORD = 'outback'     # projects containing this (e.g. Outback-MB48) are never created in Mercury
CHANGE_NUMBER = re.compile(r'(?<![A-Z])(DEV|ECO)[\s_-]*(\d{5})(?!\d)', re.IGNORECASE)
OUTPUT_SUFFIX = 'Mercury Entry Updated'

AVAILABLE = 'Available in mercury'
NOT_AVAILABLE = 'Not Available in mercury'
NO_NEED = 'No Need to Create'


class ExportError(ValueError):
    pass


def _clean(value):
    return '' if value is None else str(value).strip()


def mercury_base(item_number):
    """ASY-71940-131 -> ASY-71940-31: a three-digit last segment drops its first digit.
    Anything else (ASY-02505-21, IC-00688) is kept as-is."""
    m = THREE_DIGIT_TAIL.match(item_number)
    if m and item_number.count('-') >= 2:
        return m.group(1) + m.group(2)
    return item_number


def change_number(filename):
    """'DEV-12345.xlsx' / 'eco_12345 export.xlsx' -> 'DEV-12345' / 'ECO-12345'; None if the name has neither."""
    m = CHANGE_NUMBER.search(filename)
    return f'{m.group(1).upper()}-{m.group(2)}' if m else None


def output_filename(change):
    return f'{change} {OUTPUT_SUFFIX}.xlsx' if change else f'{OUTPUT_SUFFIX}.xlsx'


def combine(item_number, revision):
    base = mercury_base(item_number)
    return f'{base}-{revision}' if revision else base


def read_dev_export(file):
    """Return [(item_number, revision, project_name)] from the dev export, locating the header row by name.
    The project column ("Project Name F" in Agile exports) is optional."""
    try:
        ws = load_workbook(file, read_only=True, data_only=True).worksheets[0]
    except Exception as exc:
        raise ExportError(f'Could not read the dev export: {exc}') from exc

    rows = ws.iter_rows(values_only=True)
    item_col = rev_col = project_col = None
    for _, row in zip(range(HEADER_SCAN_ROWS), rows):
        headers = [_clean(c).lower() for c in row]
        if 'item number' in headers and 'revision' in headers:
            item_col, rev_col = headers.index('item number'), headers.index('revision')
            project_col = next((i for i, h in enumerate(headers) if h.startswith('project name')), None)
            break
    if item_col is None:
        raise ExportError('Dev export has no header row with "Item Number" and "Revision" columns.')

    def cell(row, col):
        return _clean(row[col]) if col is not None and col < len(row) else ''

    return [(cell(row, item_col), cell(row, rev_col), cell(row, project_col))
            for row in rows if cell(row, item_col)]


def read_mercury_export(file):
    """Every non-empty cell in the first sheet, normalised for matching."""
    try:
        ws = load_workbook(file, read_only=True, data_only=True).worksheets[0]
    except Exception as exc:
        raise ExportError(f'Could not read the Mercury export: {exc}') from exc
    return {_clean(c).upper() for row in ws.iter_rows(values_only=True) for c in row if _clean(c)}


def _is_no_need(project):
    return NO_NEED_PROJECT_WORD in project.lower()


def build_rows(dev_items, mercury_parts):
    """Only ASY/PCA parts are checked; Outback parts are marked No Need to Create whether or not they match."""
    rows = []
    for item, rev, project in dev_items:
        if not item.upper().startswith(CHECKED_PREFIXES):
            continue
        combined = combine(item, rev)
        if _is_no_need(project):
            status = NO_NEED
        elif combined.upper() in mercury_parts:
            status = AVAILABLE
        else:
            status = NOT_AVAILABLE
        rows.append({'item_number': item, 'revision': rev, 'project': project, 'combined': combined, 'status': status})
    return rows


def rows_to_xlsx(rows, change=None):
    """The change number (DEV-/ECO-) goes in a title line above the table when known."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Mercury check'
    header_row = 1
    if change:
        ws.append([change])
        ws['A1'].font = Font(bold=True, size=14)
        ws.append([])
        header_row = 3
    ws.append(['Item Number', 'Revision', 'Combined', 'Project Name', 'Status'])
    for r in rows:
        ws.append([r['item_number'], r['revision'], r['combined'], r['project'], r['status']])

    header_fill = PatternFill('solid', fgColor='5850EC')
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    status_fills = {AVAILABLE: PatternFill('solid', fgColor='E9F7EE'),
                    NOT_AVAILABLE: PatternFill('solid', fgColor='FDEDED'),
                    NO_NEED: PatternFill('solid', fgColor='F0F2F5')}
    for row in ws.iter_rows(min_row=header_row + 1, min_col=5, max_col=5):
        row[0].fill = status_fills[row[0].value]
        row[0].alignment = Alignment(horizontal='center')
    for i, width in enumerate([22, 12, 26, 24, 28], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = f'A{header_row + 1}'
    ws.auto_filter.ref = f'A{header_row}:E{ws.max_row}'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
